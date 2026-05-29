from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

def hex_to_rgb(hex_str):
    h = hex_str.lstrip('#')
    return RGBColor(*(int(h[i:i+2], 16) for i in (0, 2, 4)))

def create_rich_deck(task_id, slide_data_list, theme_color_hex="#0284c7"):
    prs = Presentation()
    prs.slide_width = Inches(13.33)  
    prs.slide_height = Inches(7.5)
    
    accent_color = hex_to_rgb(theme_color_hex)
    dark_gray = RGBColor(30, 41, 59)   
    blank_layout = prs.slide_layouts[6]

    for index, data in enumerate(slide_data_list):
        slide = prs.slides.add_slide(blank_layout)
        
        if index % 3 == 0:
            # Layout 1: Split banner
            left_banner = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(4.5), Inches(7.5))
            left_banner.fill.solid()
            left_banner.fill.fore_color.rgb = accent_color
            left_banner.line.color.rgb = accent_color

            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(3.5), Inches(3))
            tf = title_box.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = data["title"]
            tf.paragraphs[0].font.size = Pt(36)
            tf.paragraphs[0].font.bold = True
            tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

            content_box = slide.shapes.add_textbox(Inches(5.5), Inches(1.5), Inches(7.0), Inches(5))
            ctf = content_box.text_frame
            ctf.word_wrap = True
            for pt in data["points"]:
                cp = ctf.add_paragraph()
                cp.text = pt
                cp.font.size = Pt(18)
                cp.font.color.rgb = dark_gray
                cp.space_after = Pt(20)
        else:
            # Layout 2: Clean Editorial
            title_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.33), Inches(1))
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = dark_gray

            content_box = slide.shapes.add_textbox(Inches(1), Inches(3.0), Inches(10.5), Inches(3.8))
            ctf = content_box.text_frame
            ctf.word_wrap = True
            for pt in data["points"]:
                cp = ctf.add_paragraph()
                cp.text = f"→  {pt}"
                cp.font.size = Pt(16)
                cp.font.color.rgb = dark_gray
                cp.space_after = Pt(16)

    os.makedirs("./outputs", exist_ok=True)
    output_path = f"./outputs/{task_id}_rich_presentation.pptx"
    prs.save(output_path)
    return output_path